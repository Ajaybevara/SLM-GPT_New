"""Offline regression checks; never modify the user's database or graph."""
import asyncio
import io
import json
import tempfile
import unittest
from pathlib import Path
from unittest.mock import AsyncMock, patch

import numpy as np
import pymupdf
from docx import Document
from streamlit.testing.v1 import AppTest

from modules import rag_chatbot as rag
from modules.precision_ocr import generate_word_from_markdown, generate_pdf_from_markdown
from database.db_utils import register_user, verify_login


def upload(data, name):
    value = io.BytesIO(data)
    value.name = name
    return value


class ApplicationTests(unittest.TestCase):
    def test_long_password_validation(self):
        self.assertFalse(register_user('test', 'x' * 73)[0])
        self.assertEqual(verify_login('test', 'x' * 73), (False, None))

    def test_bad_upload_and_chat_failure_remain_visible(self):
        with patch('database.db_utils.initialize_extended_schema'), \
             patch.object(rag, 'get_chat_history', return_value=[]), \
             patch.object(rag, 'save_chat_message'), \
             patch.object(rag.st, 'file_uploader', return_value=[upload(b'bad PDF', 'broken.pdf')]), \
             patch.object(rag, '_async_query_graph', new_callable=AsyncMock, side_effect=RuntimeError('Model unavailable')):
            app = AppTest.from_file(str(Path(__file__).resolve().parents[1] / 'app.py'))
            app.session_state['graph_built'] = True
            app.run(timeout=30)
            next(b for b in app.button if b.label == 'Expand Graph').click().run()
            self.assertFalse(app.exception)
            self.assertTrue(any('Could not read broken.pdf' in error.value for error in app.error))
            app.chat_input[0].set_value('What is the pressure?').run(timeout=30)
            self.assertFalse(app.exception)
            self.assertTrue(any('Model unavailable' in error.value for error in app.error))

    def test_word_tables_and_paragraph_order(self):
        doc = Document()
        doc.add_paragraph('Before')
        table = doc.add_table(rows=1, cols=2)
        table.cell(0, 0).text = 'Pressure'
        table.cell(0, 1).text = '800 psi'
        doc.add_paragraph('After')
        data = io.BytesIO()
        doc.save(data)
        self.assertEqual(rag.extract_text_from_file(upload(data.getvalue(), 'report.docx')),
                         'Before\nPressure | 800 psi\nAfter')

    def test_encrypted_and_mixed_pdf(self):
        with pymupdf.open() as pdf:
            pdf.new_page().insert_text((72, 72), 'Text page from drilling report.')
            pdf.new_page()
            data = pdf.tobytes(encryption=pymupdf.PDF_ENCRYPT_AES_256,
                               owner_pw='owner', user_pw='')
        with patch.object(rag, 'vision_complete', return_value='Scanned page text') as vision:
            text = rag.extract_text_from_file(upload(data, 'mixed.pdf'))
            self.assertIn('Text page', text)
            self.assertIn('Scanned page text', text)
            vision.assert_called_once()

    def test_password_pdf_has_clear_error(self):
        with pymupdf.open() as pdf:
            pdf.new_page()
            data = pdf.tobytes(encryption=pymupdf.PDF_ENCRYPT_AES_256,
                               owner_pw='owner', user_pw='secret')
        with self.assertRaisesRegex(ValueError, 'requires a password'):
            rag.extract_text_from_file(upload(data, 'locked.pdf'))

    def test_text_bom_and_empty_text(self):
        self.assertEqual(rag.extract_text_from_file(upload(b'\xef\xbb\xbfhello', 'text.txt')), 'hello')
        self.assertEqual(rag.extract_text_from_file(upload(b'', 'text.txt')), '')

    def test_word_export_empty_cells_and_last_row(self):
        data = generate_word_from_markdown('| Field | Value |\n|---|---|\n| | 800 psi |')
        doc = Document(io.BytesIO(data))
        self.assertEqual(doc.tables[0].cell(1, 0).text, '')
        self.assertEqual(doc.tables[0].cell(1, 1).text, '800 psi')
        self.assertFalse(any('800 psi' in p.text for p in doc.paragraphs))

    def test_pdf_export(self):
        data = generate_pdf_from_markdown('# Report\n\n| Field | Value |\n|---|---|\n| Pressure | 800 psi |')
        with pymupdf.open(stream=data, filetype='pdf') as pdf:
            self.assertIn('800 psi', ''.join(p.get_text() for p in pdf))

    def test_index_and_query_in_isolated_workspace(self):
        async def vectors(texts, **kwargs):
            result = np.zeros((len(texts), rag.EMBEDDING_DIM), dtype=np.float32)
            result[:, 0] = 1
            return result
        with tempfile.TemporaryDirectory() as directory:
            module_path = Path(directory) / 'modules' / 'rag_chatbot.py'
            with patch.object(rag, '__file__', str(module_path)), \
                 patch.object(rag, 'local_embeddings', side_effect=vectors), \
                 patch.object(rag, 'llm_complete', new_callable=AsyncMock, return_value='800 psi') as llm:
                asyncio.run(rag._async_insert_text(99, 'Pressure is 800 psi.'))
                asyncio.run(rag._async_insert_text(99, 'Pressure is 800 psi.'))
                workspace = Path(directory) / 'database' / 'rag_workspace_user_99'
                self.assertTrue(rag._workspace_has_processed_documents(str(workspace)))
                self.assertEqual(len(json.loads((workspace / 'chunks.json').read_text())), 1)
                self.assertIn('800 psi', asyncio.run(rag._async_query_graph(99, 'Pressure?', 'English')))
                self.assertIn('Pressure is 800 psi.', llm.call_args.kwargs['system_prompt'])

    def test_themes_and_vault_navigation(self):
        records = [(101, 'First.pdf', 'RAG_CONTEXT', 100, '2026-10-16'),
                   (102, 'Second.pdf', 'RAG_CONTEXT', 100, '2026-10-16')]
        with patch('database.db_utils.initialize_extended_schema'), \
             patch('database.db_utils.get_user_vault', return_value=records), \
             patch('database.db_utils.get_chat_history', return_value=[]), \
             patch('database.db_utils.get_document_markdown', return_value='Pressure is 800 psi.'), \
             patch.object(rag, 'get_document_markdown', return_value='Pressure is 800 psi.'), \
             patch.object(rag, '_async_insert_text', new_callable=AsyncMock, return_value=True) as insert:
            for theme in ('dark', 'light'):
                app = AppTest.from_file(str(Path(__file__).resolve().parents[1] / 'app.py'))
                app.query_params['theme'] = theme
                app.run(timeout=30)
                self.assertFalse(app.exception)
                app.radio(key='module_navigation').set_value('Document Vault').run()
                app.checkbox(key='rag_sel_101').check()
                app.checkbox(key='rag_sel_102').check()
                app.run()
                next(b for b in app.button if 'Load 2 Documents' in b.label).click().run(timeout=30)
                self.assertFalse(app.exception)
                self.assertEqual(app.radio(key='module_navigation').value, 'Knowledge RAG')
                self.assertIn('First.pdf', insert.call_args.args[1])
                self.assertIn('Second.pdf', insert.call_args.args[1])
                app.radio(key='module_navigation').set_value('Document Vault').run()
                app.checkbox(key='temp_chat_sel_101').check()
                app.checkbox(key='temp_chat_sel_102').check()
                app.run()
                app.session_state['temp_chat_history'] = [{'role': 'user', 'content': 'Old documents'}]
                next(b for b in app.button if 'Start Temp Chat with 2' in b.label).click().run(timeout=30)
                self.assertFalse(app.exception)
                self.assertTrue(app.session_state['is_temp_rag_chat'])
                self.assertEqual(app.session_state['temp_chat_history'], [])


if __name__ == '__main__':
    unittest.main()
